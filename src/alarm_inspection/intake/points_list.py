"""Deterministic XLSX Points List normalization."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
import re
from typing import Iterable

from alarm_inspection.domain.points import decide_point, normalize_point


@dataclass(frozen=True)
class NormalizedPoint:
    address: int | None
    text: str
    accepted: bool
    reason: str
    source_row: int


def _header_score(value: object, aliases: set[str]) -> int:
    normalized = "" if value is None else re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())
    aliases = {re.sub(r"[^a-z0-9]+", "", alias.lower()) for alias in aliases}
    return 1 if normalized in aliases else 0


def _header_values(rows: list[list[object]], row_index: int) -> list[list[object]]:
    """Try a single header row, then vertically wrapped header cells."""
    row = rows[row_index]
    if row_index == 0:
        return [row]
    previous = rows[row_index - 1]
    wrapped = [
        " ".join(str(values[index]).strip() for values in (previous, row)
                 if index < len(values) and values[index] is not None)
        for index in range(max(len(previous), len(row)))
    ]
    return [row, wrapped]


def detect_columns(rows: list[list[object]]) -> tuple[int, int, int]:
    """Return header row, point column, and text column indexes."""
    point_aliases = {
        "point", "points", "point assignment", "point assignments", "point number",
        "point numbers", "point #", "point no", "point id", "address", "device address",
    }
    text_aliases = {"text", "point text", "description", "device", "device type", "device description"}
    for row_index in range(min(30, len(rows))):
        for row in _header_values(rows, row_index):
            point_columns = [i for i, value in enumerate(row) if _header_score(value, point_aliases)]
            text_columns = [i for i, value in enumerate(row) if _header_score(value, text_aliases)]
            for point_index in point_columns:
                for text_index in text_columns:
                    if point_index != text_index:
                        return row_index, point_index, text_index
    raise ValueError("Could not identify point and text columns")


def normalize_rows(rows: Iterable[list[object]]) -> list[NormalizedPoint]:
    materialized = list(rows)
    header_row, point_col, text_col = detect_columns(materialized)
    output_rows: list[dict[str, object]] = []
    orphan_text: list[str] = []
    orphan_start_row: int | None = None
    current_point_index: int | None = None
    # Repeated print headers and page furniture are structural boundaries.
    single_signature = materialized[header_row]
    try:
        detect_columns([single_signature])
        header_signature = single_signature
    except ValueError:
        header_signature = _header_values(materialized, header_row)[-1]
    repeated_headers = {
        index for index in range(header_row, len(materialized))
        if single_signature == materialized[index]
        and (header_signature == single_signature
             or header_signature in _header_values(materialized, index))
    }
    header_starts = {index - 1 for index in repeated_headers
                     if header_signature != single_signature}
    in_page_furniture = False

    def flush_orphan_text() -> None:
        nonlocal orphan_start_row
        if not orphan_text:
            return
        output_rows.append({
            "point": None,
            "text": " ".join(orphan_text),
            "source_row": orphan_start_row if orphan_start_row is not None else len(materialized),
        })
        orphan_text.clear()
        orphan_start_row = None

    for row_number, row in enumerate(materialized[header_row + 1 :], header_row + 2):
        row_index = row_number - 1
        footer = " ".join(str(value).strip() for value in row if value is not None)
        if re.search(r"\bPage\s+\d+\s+of\s+\d+\s*$", footer, re.IGNORECASE):
            in_page_furniture = True
        if row_index in repeated_headers or row_index in header_starts or in_page_furniture:
            flush_orphan_text()
            current_point_index = None
            if row_index in repeated_headers:
                in_page_furniture = False
            continue
        point_value = row[point_col] if point_col < len(row) else None
        text_value = row[text_col] if text_col < len(row) else None
        text = "" if text_value is None else str(text_value).strip()
        address = normalize_point(point_value)

        if point_value not in (None, ""):
            flush_orphan_text()
            output_rows.append({"point": point_value, "text": text, "source_row": row_number})
            current_point_index = None
            if address is not None and 1 <= address <= 255:
                current_point_index = len(output_rows) - 1
            continue

        if text:
            if current_point_index is not None:
                existing = str(output_rows[current_point_index]["text"] or "").strip()
                output_rows[current_point_index]["text"] = " ".join(
                    part for part in [existing, text] if part
                )
            else:
                if orphan_start_row is None:
                    orphan_start_row = row_number
                orphan_text.append(text)
            continue

        # Blank row creates a structural boundary for continuation text.
        flush_orphan_text()
        current_point_index = None

    flush_orphan_text()

    output: list[NormalizedPoint] = []
    for item in output_rows:
        decision = decide_point({"point": item["point"], "text": item["text"]})
        output.append(
            NormalizedPoint(
                decision.address,
                str(item["text"]),
                decision.accepted,
                decision.reason,
                int(item["source_row"]),
            )
        )
    return output


def parse_xlsx(path: str | Path) -> list[dict]:
    """Read the first usable worksheet and return normalized decisions."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    last_error: ValueError | None = None
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        try:
            normalized = normalize_rows(rows)
        except ValueError as exc:
            last_error = exc
            continue
        if normalized:
            return [asdict(item) for item in normalized]
    raise last_error or ValueError("No point rows were found in the workbook")

