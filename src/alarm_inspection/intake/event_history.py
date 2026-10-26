"""Deterministic Event History normalization."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
import re
from typing import Iterable

from alarm_inspection.domain.points import first_event_by_point, normalize_point


def _normalize_header(value: object) -> str:
    return "" if value is None else re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


_DATE_HEADERS = frozenset(_normalize_header(value) for value in (
    "date", "datetime", "time", "timestamp", "event date", "event time",
))
_POINT_HEADERS = frozenset(_normalize_header(value) for value in (
    "zone", "point", "point number", "point #", "address", "device address",
))
_STATE_HEADERS = frozenset(_normalize_header(value) for value in (
    "state", "event type", "type", "status",
))


def detect_columns(rows: list[list[object]]) -> tuple[int, int, int, int]:
    """Find the first complete header within 30 rows, taking each field's first column."""
    for row_index, row in enumerate(rows[:30]):
        date_col = point_col = state_col = None
        for column_index, value in enumerate(row):
            header = _normalize_header(value)
            if date_col is None and header in _DATE_HEADERS:
                date_col = column_index
            if point_col is None and header in _POINT_HEADERS:
                point_col = column_index
            if state_col is None and header in _STATE_HEADERS:
                state_col = column_index
            # A complete header has the old detector's maximum score. Later rows
            # cannot improve it, and its first matching columns win every tie.
            if date_col is not None and point_col is not None and state_col is not None:
                return row_index, date_col, point_col, state_col
    raise ValueError("Could not identify Date, Zone, and State columns")


def _parse_timestamp(value: object) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)

    text = str(value).strip()
    if not text:
        return None

    parsers = (
        datetime.fromisoformat,
        lambda raw: datetime.strptime(raw, "%m/%d/%Y %H:%M:%S"),
        lambda raw: datetime.strptime(raw, "%m/%d/%Y %H:%M"),
        lambda raw: datetime.strptime(raw, "%m/%d/%Y"),
        lambda raw: datetime.strptime(raw, "%Y-%m-%d %H:%M:%S"),
        lambda raw: datetime.strptime(raw, "%Y-%m-%d %H:%M"),
        lambda raw: datetime.strptime(raw, "%Y-%m-%d"),
    )
    for parser in parsers:
        try:
            return parser(text)
        except ValueError:
            continue
    return None


def normalize_rows(rows: Iterable[list[object]]) -> dict[int, datetime]:
    materialized = list(rows)
    header_row, date_col, zone_col, state_col = detect_columns(materialized)
    events: list[dict[str, object]] = []

    for row in materialized[header_row + 1 :]:
        date_value = row[date_col] if date_col < len(row) else None
        zone_value = row[zone_col] if zone_col < len(row) else None
        state_value = row[state_col] if state_col < len(row) else None

        timestamp = _parse_timestamp(date_value)
        if timestamp is None:
            continue

        state = "" if state_value is None else str(state_value).strip().upper()
        if state not in {"A", "T"}:
            continue

        point = normalize_point(zone_value)
        if point is None:
            continue

        events.append({"event_type": state, "point": point, "timestamp": timestamp})

    earliest = first_event_by_point(events)
    return {int(point): timestamp for point, timestamp in earliest.items() if isinstance(timestamp, datetime)}


def parse_xlsx(path: str | Path) -> dict[int, datetime]:
    """Read the first usable worksheet and return earliest qualifying event per point."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    last_error: ValueError | None = None
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        try:
            normalized = normalize_rows(rows)
        except ValueError as exc:
            last_error = exc
            continue
        if normalized:
            return normalized
    raise last_error or ValueError("No qualifying Event History rows were found in the workbook")